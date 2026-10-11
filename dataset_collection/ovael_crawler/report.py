"""Generate dataset_documentation.docx from the actual output files."""

from __future__ import annotations

import csv
import json
import statistics
from collections import Counter
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

from .reference import TrackerReference, utcnow
from .schema import FIELDS, category_columns

TITLE = "AI-Based Website Behavioral Analysis and Tracker Detection System"


def _table(doc, header, rows, widths=None, font=8):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = str(h)
        for r in cell.paragraphs[0].runs:
            r.bold = True
            r.font.size = Pt(font)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = "" if v is None else str(v)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(font)
    doc.add_paragraph()
    return t


def _bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def _num(v):
    if v in ("", None):
        return None
    try:
        return float(v)
    except ValueError:
        return None


def _fmt(x):
    if x is None:
        return "-"
    return f"{x:.3f}".rstrip("0").rstrip(".") if isinstance(x, float) else str(x)


def build_report(out: Path) -> Path:
    manifest = json.loads((out / "crawl_manifest.json").read_text())
    recs = [json.loads(l) for l in (out / "website_observations.jsonl").read_text().splitlines() if l.strip()] \
        if (out / "website_observations.jsonl").exists() else []
    with open(out / "website_features.csv", newline="") as f:
        rows = list(csv.DictReader(f))
    ref = TrackerReference()
    st = manifest["status"]
    cfg = manifest.get("crawl_config", {})
    src = manifest.get("source", {})
    is_smoke = manifest.get("run_type") == "smoke"
    preliminary = not st.get("complete") or is_smoke

    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(TITLE, 0)
    doc.add_paragraph("Website Behavioral Dataset - Technical Documentation").runs[0].bold = True
    p = doc.add_paragraph()
    banner = ("PRELIMINARY - SMOKE-TEST OUTPUT (not the research dataset)" if is_smoke else
              "PRELIMINARY - COLLECTION INCOMPLETE" if preliminary else "FINAL - COLLECTION COMPLETE")
    r = p.add_run(banner)
    r.bold = True
    r.font.color.rgb = RGBColor(0xB0, 0x00, 0x00) if preliminary else RGBColor(0x00, 0x70, 0x30)
    doc.add_paragraph(f"Generated {utcnow()} from {out.name}/ (website_observations.jsonl, website_features.csv, "
                      f"crawl_manifest.json). All statistics below are computed from these files; nothing is estimated.")

    doc.add_heading("1. Objective", 1)
    doc.add_paragraph(
        "Collect real, browser-observed privacy behaviour (network requests, third-party and tracker-list "
        "matched requests, cookies, Web Storage use, fingerprinting-related API calls, redirects and iframes) "
        "for a fixed set of popular websites, as a behavioural feature dataset for exploratory analysis and an "
        "initial unsupervised anomaly-detection baseline for the Ovael tracker/anomaly detection system.")

    doc.add_heading("2. Collection methodology", 1)
    _bullets(doc, [
        f"Browser: Playwright {manifest.get('environment', {}).get('playwright', '?')} driving Chromium "
        f"{manifest.get('environment', {}).get('browser_version', '?')}, headless={cfg.get('headless')}, default "
        f"Playwright user agent, viewport {cfg.get('viewport_width')}x{cfg.get('viewport_height')}, locale {cfg.get('locale')}.",
        "Isolation: a new browser context (empty cookie jar and storage, service workers blocked, downloads "
        "refused) for every navigation attempt; the record keeps the observations of the final attempt only.",
        f"URL strategy: https://<domain>/ first; if the host is unreachable (DNS failure, refused connection, TLS "
        f"error) then https://www.<domain>/, then http://<domain>/. Transient failures (timeouts, resets, HTTP 429/503) "
        f"are retried up to {cfg.get('max_retries')} times with exponential backoff (base {cfg.get('backoff_base_s')} s, "
        f"Retry-After honoured up to {cfg.get('max_retry_after_s')} s).",
        f"Timing: page.goto waits for DOMContentLoaded (limit {cfg.get('navigation_timeout_s')} s), then up to "
        f"{cfg.get('load_event_timeout_s')} s for the load event, then a fixed settling period of {cfg.get('settle_s')} s. "
        f"Observations (cookies, DOM iframes, instrumentation counters) are read once, at the end of the settling "
        f"period. Hard cap per website: {cfg.get('overall_timeout_s')} s. Concurrency: {cfg.get('workers')} workers.",
        "No interaction: no clicks, scrolling, form submission, consent-banner handling, login, or CAPTCHA solving. "
        "Observations therefore reflect the pre-consent landing-page state.",
        "Network evidence: Playwright request/response/requestfailed events on the context (all frames and popups).",
        "Instrumentation: an init script injected before any page script in every frame wraps selected Canvas, WebGL, "
        "Audio, font, battery, media-device, WebRTC, permissions, Web Storage, document.cookie, IndexedDB, Cache "
        "Storage and StorageManager APIs and counts calls. Counts from documents that were replaced or removed are "
        "reported through a Playwright binding; live documents are read directly at the end.",
    ])

    doc.add_heading("3. Source list and selection", 1)
    if is_smoke:
        doc.add_paragraph(f"Smoke-test list: {src.get('source_file')} (fixed list of well-known sites plus "
                          "deliberate failure / duplicate / unsafe-redirect test entries).")
    else:
        _table(doc, ["Item", "Value"], [
            ["Source", "Tranco list (https://tranco-list.eu)"],
            ["List ID", src.get("list_id")],
            ["Identity verified", f"{src.get('verified')} ({src.get('verification_method')})"],
            ["List date / created", f"{src.get('list_date')} / {src.get('list_created_on')}"],
            ["List configuration", json.dumps(src.get("list_configuration"))],
            ["Permalink", src.get("list_permalink")],
            ["Local file", f"{src.get('source_file')} (sha256 {src.get('source_file_sha256')})"],
            ["Download date (local file mtime, UTC)", src.get("source_file_mtime_utc")],
        ])
    sel = manifest["selection"]
    doc.add_paragraph("Selection method: " + sel["method"])
    doc.add_paragraph(f"Selected: {sel['selected_count']} websites; skipped while selecting: {sel['skipped_count']} "
                      f"entries ({dict(Counter(s['reason'].split(':')[0] for s in sel['skipped']))}). "
                      f"Malware host check: {sel['malware_host_check'] if isinstance(sel['malware_host_check'], str) else sel['malware_host_check'].get('source') + ', retrieved ' + str(sel['malware_host_check'].get('retrieved_at'))}.")

    doc.add_heading("4. Collection status", 1)
    _table(doc, ["Selected", "Attempted", "Success", "HTTP error (observed)", "Failed", "Blocked (unsafe)", "Pending", "Complete"],
           [[st["selected"], st["attempted"], st["success"], st["http_error"], st["failed"], st["blocked_unsafe"],
             st["pending"], st["complete"]]])
    doc.add_paragraph("Attempted = websites with a final record. 'HTTP error' pages (main document status >= 400, often bot "
                      "challenges or geo-blocks) were loaded and observed; their features describe the error page.")
    if recs:
        times = sorted(r["crawl"]["crawl_timestamp"] for r in recs)
        doc.add_paragraph(f"Collection window (UTC): {times[0]} to {max(r['crawl']['finished_at'] for r in recs)}.")
    fails = Counter(r["crawl"]["navigation_outcome"] for r in recs if r["crawl"]["status"] in ("failed", "blocked_unsafe"))
    if fails:
        doc.add_paragraph("Failure classes:")
        _table(doc, ["Failure class", "Websites"], fails.most_common())
    http_err = Counter(str(r["final"].get("http_status")) for r in recs if r["crawl"]["status"] == "http_error")
    if http_err:
        doc.add_paragraph("Main-document HTTP status of 'http_error' visits:")
        _table(doc, ["HTTP status", "Websites"], http_err.most_common())
    outcomes = Counter(r["crawl"]["navigation_outcome"] for r in recs if r["crawl"]["status"] in ("success", "http_error"))
    if outcomes:
        _table(doc, ["Navigation outcome (observed pages)", "Websites"], outcomes.most_common())
    runs = manifest.get("runs", [])
    if runs:
        doc.add_paragraph("Crawler runs (each resume is a run):")
        _table(doc, ["Started", "Finished", "Scheduled", "Stop reason"],
               [[x.get("started_at"), x.get("finished_at"), x.get("scheduled"), x.get("stop_reason")] for x in runs])

    doc.add_heading("5. Configuration", 1)
    _table(doc, ["Parameter", "Value"], [[k, json.dumps(v) if isinstance(v, list) else v] for k, v in cfg.items()])
    env = manifest.get("environment", {})
    if env:
        _table(doc, ["Environment", "Value"], [[k, v] for k, v in env.items()])

    doc.add_heading("6. Tracker reference list", 1)
    v = ref.version
    _table(doc, ["Item", "Value"], [
        ["Source", v["source"]], ["URL", v["url"]], ["Retrieved", v["retrieved_at"]],
        ["HTTP Last-Modified", v["http_last_modified"]], ["SHA-256", v["sha256"]],
        ["Built with (Tracker Radar commit)", json.dumps(v["built_with"])], ["Tracker entries", v["tracker_entries"]],
        ["Categories", ", ".join(ref.categories)],
    ])
    doc.add_paragraph(
        "Matching rule: a request hostname matches when it, or a parent domain, is a key of the 'trackers' object. "
        "Only third-party requests (registrable domain differs from the final page's) are counted as tracker-matched "
        "requests. The list is DuckDuckGo's curated blocking subset (~1,000 tracker root domains with owner, category, "
        "prevalence and a 0-3 fingerprinting score); a hostname absent from it is recorded as 'not_in_reference' - "
        "unknown, never 'benign'. Matching is domain-level; DuckDuckGo's per-path rules and exceptions are not applied, "
        "so matches can include requests DuckDuckGo would allow (default_action and same_entity_as_site are recorded).")

    doc.add_heading("7. Summary statistics (from collected data)", 1)
    numeric = [f for f in FIELDS if f.dtype in ("integer", "float", "boolean") and f.group != "labels"
               and f.name not in ("tranco_rank", "http_status")]
    srows = []
    for fdef in numeric:
        vals = [_num(r[fdef.name]) for r in rows]
        present = [x for x in vals if x is not None]
        if present:
            srows.append([fdef.name, len(present), len(vals) - len(present), _fmt(statistics.mean(present)),
                          _fmt(statistics.median(present)), _fmt(min(present)), _fmt(max(present)),
                          f"{100 * sum(1 for x in present if x == 0) / len(present):.0f}%"])
        else:
            srows.append([fdef.name, 0, len(vals), "-", "-", "-", "-", "-"])
    _table(doc, ["Feature", "n", "missing", "mean", "median", "min", "max", "% zero"], srows, font=7)
    site_tr = Counter()
    for r in recs:
        for d in {hst["matched_tracker_domain"] for hst in (r.get("tracker_summary") or {}).get("matched_hosts", [])}:
            site_tr[d] += 1
    n_obs = sum(1 for r in recs if r["crawl"]["status"] in ("success", "http_error"))
    if site_tr:
        doc.add_paragraph(f"Most widespread tracker-list domains (number of the {n_obs} observed websites on which each appeared as third party):")
        owners = {}
        for r in recs:
            for hst in (r.get("tracker_summary") or {}).get("matched_hosts", []):
                owners[hst["matched_tracker_domain"]] = hst.get("owned_by") or hst.get("owner_name")
        _table(doc, ["Tracker domain", "Owner", "Websites"], [[d, owners.get(d), c] for d, c in site_tr.most_common(20)])
    cats = Counter()
    for r in recs:
        for c in (r.get("tracker_summary") or {}).get("requests_by_category", {}):
            cats[c] += 1
    if cats:
        doc.add_paragraph("Tracker categories (websites with >= 1 matched third-party request in the category):")
        _table(doc, ["Category", "Websites"], cats.most_common())
    if n_obs:
        def share(pred):
            return f"{100 * sum(1 for r in rows if r['crawl_status'] in ('success', 'http_error') and pred(r)) / n_obs:.1f}%"
        _table(doc, ["Indicator (observed websites)", "Share"], [
            ["At least one tracker-list match", share(lambda r: r["has_tracker_reference_match"] == "1")],
            ["At least one third-party cookie", share(lambda r: (_num(r["third_party_cookies"]) or 0) > 0)],
            ["Canvas read-back call (toDataURL/toBlob/getImageData)", share(lambda r: (_num(r["canvas_readback_calls"]) or 0) > 0)],
            ["WebGL unmasked vendor/renderer query", share(lambda r: (_num(r["webgl_unmasked_info_queries"]) or 0) > 0)],
            ["OfflineAudioContext created", share(lambda r: (_num(r["offline_audio_contexts"]) or 0) > 0)],
            ["localStorage write", share(lambda r: (_num(r["localstorage_write_count"]) or 0) > 0)],
            ["Hidden iframe", share(lambda r: (_num(r["hidden_iframe_count"]) or 0) > 0)],
            ["Cross-domain redirect", share(lambda r: r["cross_domain_redirect"] == "1")],
        ])
        doc.add_paragraph("Fingerprinting-related API calls are observed indicators only; calling these APIs does not by "
                          "itself establish that a website fingerprints visitors.")

    doc.add_heading("8. Feature definitions", 1)
    doc.add_paragraph("website_features.csv columns (one row per selected website that has been attempted). Tracker "
                      f"category columns ({len(ref.categories)}) follow 'third_party_hostnames_not_in_reference': "
                      + ", ".join(category_columns(ref.categories)) +
                      " - each is the number of tracker-matched third-party requests whose reference entry lists that category "
                      "(a request can count in several categories).")
    _table(doc, ["Field", "Type", "Group", "Definition", "Missing value"],
           [[f.name, f.dtype, f.group, f.description, f.missing] for f in FIELDS], font=7)

    doc.add_heading("9. Output files and data schema", 1)
    _table(doc, ["File", "Content"], [
        ["website_features.csv", "One row per attempted website: metadata, crawl status, numerical features, label placeholders."],
        ["website_observations.jsonl", "One JSON object per attempted website: attempts and errors, final URL, redirect chain, "
         "request/tracker/cookie summaries, third-party host list with reference match status, frames and DOM iframes, "
         "instrumentation counters and evidence events, safety-proxy blocks, label fields. This file is the commit log "
         "from which both CSVs are regenerated."],
        ["request_observations.jsonl", "One JSON object per observed request: sanitised URL (no query/fragment), hostname, "
         "registrable domain, party, resource type, method, frame kind, redirect link, status, failure text, content type, "
         "tracker match. No headers or bodies."],
        ["cookie_metadata.jsonl", "One JSON object per cookie: name, domain, path, expiry, Secure, HttpOnly, SameSite, party. No values."],
        ["tracker_matches.csv", "One row per (website, matched third-party hostname) with reference metadata."],
        ["crawl_manifest.json", "Selection, source list identity, reference version, configuration, environment, runs, status."],
        ["crawl.log", "Structured JSON log of the crawl."],
    ])
    observed = [r for r in rows if r["crawl_status"] in ("success", "http_error")]
    failed = [r for r in rows if r["crawl_status"] not in ("success", "http_error")]
    sample_cols = ["website_id", "original_domain", "final_url", "crawl_status", "http_status", "total_requests",
                   "third_party_requests", "tracker_matched_requests", "total_cookies", "third_party_cookies",
                   "canvas_api_calls", "webgl_api_calls", "localstorage_write_count", "iframe_count",
                   "hidden_iframe_count", "navigation_duration_seconds"]
    samples = observed[:3] + failed[:1]
    if samples:
        doc.add_paragraph("Sample records (selected columns, verbatim from website_features.csv):")
        _table(doc, ["Field"] + [s["website_id"] for s in samples],
               [[c] + [s.get(c, "") for s in samples] for c in sample_cols + ["failure_reason"]], font=7)

    doc.add_heading("10. Missing-data handling", 1)
    _bullets(doc, [
        "Empty cell = not measured. 0 = measured and none observed. Unknown values are never replaced by 0.",
        "Failed and blocked visits keep metadata, status and failure reason; all observation features are empty.",
        "Instrumentation-derived fields (storage, canvas, WebGL, audio, other APIs) are empty when the main-frame "
        "instrumentation snapshot could not be read (instrumentation_status = main_frame_unavailable).",
        "dom_iframe_count/hidden_iframe_count are empty if the top document could not be evaluated.",
        "tracker_request_ratio is empty when total_requests is 0; max_cookie_lifetime_days is empty without persistent cookies.",
        "A tracker-list miss is 'not_in_reference' (unknown), not benign.",
        "An empty cookie list means no cookies were present in the context at observation time, not that the site never uses cookies.",
    ])

    doc.add_heading("11. Limitations and potential measurement bias", 1)
    _bullets(doc, [
        "Single visit, landing page only, from one vantage point (one network/country), one point in time; content, "
        "ads and trackers vary between visits, regions and over time.",
        "Headless Chromium with the default Playwright user agent is detectable; some sites serve bot challenges, block "
        "pages or reduced content (visible as http_error, or as 'success' pages with very low dom_element_count / "
        "visible_text_length and few requests). No evasion was attempted; such pages are kept and should be filtered or "
        "flagged during analysis.",
        "No consent interaction: many EU-facing sites load fewer trackers before consent; counts are pre-consent behaviour.",
        "Tranco top domains include infrastructure/CDN/API domains without a browsable homepage; they appear as failures or "
        "minimal pages. Popularity-ranked selection over-represents large organisations.",
        "The fixed settling period truncates behaviour that starts later (lazy loading, timers, scroll-triggered scripts).",
        "Tracker reference is a curated ~1,000-domain list; coverage of small or regional trackers is incomplete, and "
        "domain-level matching ignores DuckDuckGo's path rules and same-entity exemptions (recorded separately).",
        "First/third-party attribution uses the registrable domain (PSL); different domains of one organisation "
        "(e.g. a site and its CDN) count as third party. same_entity_as_site uses the reference's entity map when available.",
        "Instrumentation does not see: Web Workers / Service Workers (no init script; service workers blocked), direct "
        "property-style storage access (localStorage.foo = x), calls made before injection in rare cases, or "
        "documents removed before their snapshot was reported. Wrapped functions are detectable by scripts that "
        "inspect them, which could change site behaviour.",
        "Cookies are read once at the end; cookies set and deleted earlier are missed. Partitioned (CHIPS) cookies are "
        "reported as Playwright exposes them.",
        "Hidden-iframe criterion only examines iframes in the top-level document; nested iframes are counted in "
        "iframe_count but not classified for visibility.",
        "Redirect counts include HTTP redirects of the main document; client-side changes are counted separately and "
        "include history-API URL changes, not only JavaScript/meta redirects.",
    ])

    doc.add_heading("12. Privacy and security safeguards", 1)
    _bullets(doc, [
        "Every connection from the browser goes through a local validating proxy that resolves DNS itself, refuses any "
        "destination with a non-public address (loopback, private, link-local, CGNAT, multicast, reserved) and connects to "
        "the validated address. This applies to redirects, subresources, iframes and WebSockets. Blocked connections "
        "are recorded (blocked_unsafe_destinations).",
        "QUIC disabled and WebRTC restricted to proxied UDP so traffic cannot bypass the proxy.",
        "Downloads refused, service workers blocked, no forms submitted, no logins, no CAPTCHA or access-control bypass.",
        "Selected domains were checked against the URLhaus malware host list when available.",
        "Not stored: cookie values, storage values, request/response headers and bodies, query strings, URL fragments, "
        "canvas pixels or drawn text. URL paths are truncated to 200 characters.",
        "Cookie values (held in memory only during a visit) are used to scrub every saved string: any occurrence of a "
        "cookie value in a URL path, cookie name, title or frame URL is replaced by <redacted>. Alphanumeric tokens of "
        "8+ characters containing a digit inside cookie names are also redacted (some sites embed identifiers in names).",
        "The smoke test verifies this: real cookie values are written to a temporary file outside the output directory, "
        "every output file is searched for them, and the temporary file is deleted.",
        "Polite crawling: one visit per site, 1 s pause per worker between sites, bounded retries with backoff, Retry-After "
        "honoured, run stops after repeated server errors or loss of connectivity.",
    ])

    doc.add_heading("13. Labelling methodology and limitations", 1)
    doc.add_paragraph(
        "No ground-truth labels were assigned. security_label, manual_review_label, label_source, label_confidence and "
        "labeled_at are present but empty. Tracker-list matches (has_tracker_reference_match, tracker_* features) are "
        "observations from a reference list, not maliciousness labels: using a website does not make it malicious, and "
        "Tranco sites are popular, mostly legitimate websites. This collection is therefore a behavioural feature dataset "
        "suitable for exploratory analysis and unsupervised anomaly detection - not a labelled training set. If tracker "
        "features are later used as labels, they must not also be used as predictive features in any evaluation that "
        "claims independent predictive performance.")

    doc.add_heading("14. Model training and validation status", 1)
    doc.add_paragraph("No model has been trained or validated on this dataset. No accuracy, precision, recall or other "
                      "model metrics exist for it. Validated classification performance requires independently "
                      "obtained ground-truth labels.")
    path = out / "dataset_documentation.docx"
    doc.save(str(path))
    return path
